"""Extract text from Markdown, DOCX, PDF, and image documents."""

from pathlib import Path
from typing import Any, Dict, Tuple


SUPPORTED_DOCUMENT_EXTENSIONS = {".md", ".markdown", ".docx", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}
OCR_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}


def _plain_document_text(text: str, source_name: str) -> str:
    """Give plain extracted files a heading so the generic chunker can split them."""
    cleaned = text.replace("\x00", "").strip()
    if not cleaned:
        raise ValueError(f"{source_name}: không trích xuất được nội dung văn bản")
    return f"### Nội dung tài liệu: {source_name}\n{cleaned}"


def load_document(path: Path) -> Tuple[Dict[str, Any], str]:
    """Return loader metadata and Markdown-like text for a supported document."""
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise ValueError(f"{path.name}: định dạng chưa được hỗ trợ")

    if suffix in {".md", ".markdown"}:
        return _load_markdown(path)

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Thiếu python-docx. Chạy: pip install python-docx") from exc
        document = Document(str(path))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return {"document_type": "docx"}, _plain_document_text(text, path.name)

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("Thiếu pypdf. Chạy: pip install pypdf") from exc
        pages = []
        for page_number, page in enumerate(PdfReader(str(path)).pages, 1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"### Trang {page_number}\n{page_text.strip()}")
        if pages:
            return {"document_type": "pdf", "page_count": len(pages)}, "\n\n".join(pages)
        try:
            import fitz
            from PIL import Image
            import pytesseract
        except ImportError as exc:
            raise RuntimeError(
                f"{path.name}: PDF scan cần PyMuPDF, Pillow và pytesseract để OCR."
            ) from exc
        ocr_pages = []
        for page_number, page in enumerate(fitz.open(str(path)), 1):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
            text = pytesseract.image_to_string(image, lang="vie+eng").strip()
            if text:
                ocr_pages.append(f"### Trang {page_number}\n{text}")
        if not ocr_pages:
            raise RuntimeError(
                f"{path.name}: không trích xuất/OCR được nội dung PDF."
            )
        return {"document_type": "pdf_ocr", "page_count": len(ocr_pages)}, "\n\n".join(ocr_pages)

    try:
        from PIL import Image
        import pytesseract
    except ImportError as exc:
        raise RuntimeError(
            "Thiếu Pillow/pytesseract. Chạy: pip install Pillow pytesseract "
            "và cài Tesseract OCR trên Windows."
        ) from exc
    try:
        text = pytesseract.image_to_string(Image.open(path), lang="vie+eng")
    except Exception as exc:
        raise RuntimeError(
            f"{path.name}: OCR thất bại. Kiểm tra Tesseract đã cài và có trong PATH."
        ) from exc
    return {"document_type": "image_ocr", "ocr_language": "vie+eng"}, _plain_document_text(text, path.name)


def _load_markdown(path: Path) -> Tuple[Dict[str, Any], str]:
    text = path.read_text(encoding="utf-8")
    if text.startswith("---"):
        parts = text.split("---", 2)
        if len(parts) == 3:
            import yaml
            return dict(yaml.safe_load(parts[1]) or {}), parts[2].lstrip()
    return {"document_type": "markdown"}, text
